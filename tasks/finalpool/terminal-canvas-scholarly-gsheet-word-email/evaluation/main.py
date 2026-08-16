import argparse
import json
import os
import sys
import psycopg2

DB_CONFIG = {
    "host": os.environ.get("PGHOST", "localhost"),
    "port": int(os.environ.get("PGPORT", "5432")),
    "dbname": os.environ.get("PGDATABASE", "toolathlon_gym"),
    "user": "eigent",
    "password": "camel",
}

PASS_COUNT = 0
FAIL_COUNT = 0


def check(name, condition, detail=""):
    global PASS_COUNT, FAIL_COUNT
    if condition:
        PASS_COUNT += 1
        print(f"  [PASS] {name}")
    else:
        FAIL_COUNT += 1
        print(f"  [FAIL] {name}: {str(detail)[:200]}")


def num_close(a, b, tol=5.0):
    try:
        return abs(float(a) - float(b)) <= tol
    except:
        return False


_FALLBACK_FACULTY = ["Dr. Sarah Chen", "Dr. James Okafor", "Dr. Maria Gonzalez", "Dr. Raj Patel"]
_FALLBACK_DEPARTMENTS = ["Biochemistry", "Bioinformatics"]


def _get_faculty_from_roster():
    """Read faculty_roster.csv to get faculty names and departments dynamically."""
    try:
        import csv as _csv
        roster_path = os.path.join(
            os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
            "initial_workspace",
            "faculty_roster.csv",
        )
        faculty = []
        departments = set()
        with open(roster_path) as f:
            reader = _csv.DictReader(f)
            for row in reader:
                name = row.get("faculty_name", row.get("name", "")).strip()
                dept = row.get("department", "").strip()
                if name:
                    faculty.append(name)
                if dept:
                    departments.add(dept)
        if faculty:
            return faculty, sorted(departments)
        return _FALLBACK_FACULTY, _FALLBACK_DEPARTMENTS
    except Exception:
        return _FALLBACK_FACULTY, _FALLBACK_DEPARTMENTS


def check_xlsx_content(workspace):
    """Check Alignment_Summary.xlsx has valid content."""
    print("\n=== Checking XLSX Content ===")
    try:
        import openpyxl
    except ImportError:
        check("openpyxl available", False, "Cannot import openpyxl")
        return False

    xlsx_path = os.path.join(workspace, "Alignment_Summary.xlsx")
    if not os.path.isfile(xlsx_path):
        check("Alignment_Summary.xlsx exists", False, f"Not found: {xlsx_path}")
        return False
    check("Alignment_Summary.xlsx exists", True)

    try:
        wb = openpyxl.load_workbook(xlsx_path, data_only=True)
        check("XLSX has at least one sheet", len(wb.worksheets) >= 1,
              f"Found {len(wb.worksheets)} sheets")
        all_ok = True
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(values_only=True))
            has_data = len(rows) >= 2
            check(f"XLSX sheet '{ws.title}' has data rows", has_data,
                  f"Only {len(rows)} rows")
            if not has_data:
                all_ok = False
        wb.close()
        return all_ok
    except Exception as e:
        check("XLSX readable", False, str(e))
        return False


FACULTY, DEPARTMENTS = _get_faculty_from_roster()
FACULTY_COUNT = len(FACULTY)

# Noise faculty/departments that should NOT appear in outputs
NOISE_DEPARTMENTS = ["Computer Science", "Mathematics", "Physics", "History"]


def check_reverse_validation():
    print("\n=== Reverse Validation ===")
    conn = psycopg2.connect(**DB_CONFIG)
    cur = conn.cursor()
    try:
        # Check GSheet does not contain irrelevant/noise faculty
        cur.execute("""
            SELECT id FROM gsheet.spreadsheets
            WHERE title ILIKE '%%Research_Teaching_Alignment%%'
               OR title ILIKE '%%Research%%Teaching%%Alignment%%'
            LIMIT 1
        """)
        row = cur.fetchone()
        if row:
            ss_id = row[0]
            cur.execute("""
                SELECT value FROM gsheet.cells
                WHERE spreadsheet_id = %s AND row_index > 0
            """, (ss_id,))
            all_values = " ".join([r[0] for r in cur.fetchall() if r[0]]).lower()
            # Only 4 faculty should be present; check no noise departments
            for dept in NOISE_DEPARTMENTS:
                check(f"GSheet does not contain noise department '{dept}'",
                      dept.lower() not in all_values,
                      f"Found '{dept}' in GSheet data")
            # Check no more than 4 distinct faculty rows
            cur.execute("""
                SELECT COUNT(DISTINCT row_index) FROM gsheet.cells
                WHERE spreadsheet_id = %s AND row_index > 0
            """, (ss_id,))
            data_rows = cur.fetchone()[0]
            check("GSheet has no extra noise rows beyond 4 faculty",
                  data_rows <= 4, f"Found {data_rows} data rows")

        # Check no emails sent to wrong recipients
        noise_emails = ["newsletter@university.edu", "all-faculty@university.edu",
                        "registrar@university.edu"]
        for addr in noise_emails:
            cur.execute(
                "SELECT COUNT(*) FROM email.messages WHERE to_addr::text ILIKE %s",
                (f"%{addr}%",),
            )
            cnt = cur.fetchone()[0]
            check(f"No email sent to noise recipient {addr}", cnt == 0,
                  f"Found {cnt} emails to {addr}")
    except Exception as e:
        check("Reverse validation", False, str(e))
    finally:
        cur.close()
        conn.close()


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--agent_workspace", required=False, default=".")
    parser.add_argument("--groundtruth_workspace", required=False, default=".")
    parser.add_argument("--launch_time", required=False)
    parser.add_argument("--res_log_file", required=False)
    args = parser.parse_args()

    ws = args.agent_workspace

    conn = psycopg2.connect(**DB_CONFIG)
    cur = conn.cursor()

    # 1. Check course_keywords.json exists and has content
    kw_path = os.path.join(ws, "course_keywords.json")
    check("course_keywords.json exists", os.path.exists(kw_path), "File not found")
    if os.path.exists(kw_path):
        with open(kw_path) as f:
            kw_data = json.load(f)
        check("course_keywords has entries", len(kw_data) > 0, f"Got {len(kw_data)} entries")
    else:
        kw_data = {}

    # 2. Check alignment_scores.json exists and has correct faculty entries
    as_path = os.path.join(ws, "alignment_scores.json")
    check("alignment_scores.json exists", os.path.exists(as_path), "File not found")
    alignment_score_values = []
    if os.path.exists(as_path):
        with open(as_path) as f:
            as_data = json.load(f)

        def _extract_score(item):
            if isinstance(item, (int, float)):
                return float(item)
            if isinstance(item, dict):
                for key in ("alignment_score", "score", "alignment_score_percent", "alignmentScore", "percentage", "percent"):
                    if key in item:
                        try:
                            return float(item[key])
                        except (TypeError, ValueError):
                            return None
            return None

        # Could be list or dict
        if isinstance(as_data, list):
            check(f"alignment_scores has {FACULTY_COUNT} entries", len(as_data) == FACULTY_COUNT, f"Got {len(as_data)}")
            names_in_scores = [e.get("faculty_name", e.get("name", "")) for e in as_data if isinstance(e, dict)]
            alignment_score_values = [_extract_score(e) for e in as_data]
        elif isinstance(as_data, dict):
            check(f"alignment_scores has {FACULTY_COUNT} entries", len(as_data) == FACULTY_COUNT, f"Got {len(as_data)}")
            names_in_scores = list(as_data.keys())
            alignment_score_values = [_extract_score(v) for v in as_data.values()]
        else:
            names_in_scores = []
            check(f"alignment_scores has {FACULTY_COUNT} entries", False, "Unexpected format")

        # Check each faculty name appears
        for fac in FACULTY:
            found = any(fac.lower() in n.lower() for n in names_in_scores)
            check(f"alignment_scores contains {fac}", found, f"Names: {names_in_scores}")

        # Validate score values are numeric percentages in [0, 100]
        non_numeric = [v for v in alignment_score_values if v is None]
        check("alignment_scores values are numeric",
              len(non_numeric) == 0 and len(alignment_score_values) == FACULTY_COUNT,
              f"Non-numeric or missing: {non_numeric}, got {len(alignment_score_values)} values")
        out_of_range = [v for v in alignment_score_values if v is not None and not (0.0 <= v <= 100.0)]
        check("alignment_scores values within [0, 100] percentage range",
              len(out_of_range) == 0,
              f"Out of range: {out_of_range}")
    else:
        as_data = {}

    # 3. Check summary_stats.json
    ss_path = os.path.join(ws, "summary_stats.json")
    check("summary_stats.json exists", os.path.exists(ss_path), "File not found")
    if os.path.exists(ss_path):
        with open(ss_path) as f:
            ss_data = json.load(f)
        # Required keys per task.md: total faculty, average alignment score, max, min, count below 30%
        for key in ["total_faculty", "average_alignment_score", "max_score", "min_score", "below_30"]:
            found = any(key.replace("_", "") in k.replace("_", "").replace(" ", "").lower() for k in ss_data.keys()) or key in ss_data
            if key == "below_30":
                # accept variations like "below_30_percent", "count_below_30",
                # "faculty_below_30", "below_threshold_30" but require an
                # explicit '30' (or 0.3) marker so it doesn't fuzzy-match
                # unrelated keys.
                found = found or any(
                    "30" in k or "0.3" in k for k in ss_data.keys()
                )
            check(f"summary_stats has {key}", found, f"Keys: {list(ss_data.keys())}")
        # Validate total_faculty value matches actual count
        for k, v in ss_data.items():
            if "total" in k.lower() and "facult" in k.lower():
                check(f"summary_stats total_faculty == {FACULTY_COUNT}",
                      v == FACULTY_COUNT, f"Got {v}")
                break

    # 4. Check Word document exists
    doc_path = os.path.join(ws, "Research_Teaching_Report.docx")
    check("Research_Teaching_Report.docx exists", os.path.exists(doc_path), "File not found")
    if os.path.exists(doc_path):
        try:
            from docx import Document
            doc = Document(doc_path)
            full_text = "\n".join([p.text for p in doc.paragraphs])
            text_lower = full_text.lower()

            check("Report has Executive Summary", "executive summary" in text_lower, "Section not found")
            check("Report mentions alignment score", "alignment" in text_lower and "score" in text_lower, "Not found")
            for fac in FACULTY:
                check(f"Report mentions {fac}", fac.lower() in text_lower, "Not found")
            for dept in DEPARTMENTS:
                check(f"Report mentions {dept} dept", dept.lower() in text_lower, "Not found")
            check("Report has recommendations section", "recommend" in text_lower, "Not found")
            check("Report mentions review needed concept", "review" in text_lower, "Not found")
        except Exception as e:
            check("Report content readable", False, str(e))

    # 5. Check Google Sheet
    cur.execute("SELECT id FROM gsheet.spreadsheets WHERE title ILIKE '%Research_Teaching_Alignment%' OR title ILIKE '%Research%Teaching%Alignment%' LIMIT 1")
    row = cur.fetchone()
    check("GSheet spreadsheet exists", row is not None, "No matching spreadsheet found")
    if row:
        ss_id = row[0]
        cur.execute("SELECT id FROM gsheet.sheets WHERE spreadsheet_id = %s AND title ILIKE '%%Alignment%%Matrix%%' LIMIT 1", (ss_id,))
        sheet_row = cur.fetchone()
        check("GSheet has Alignment Matrix sheet", sheet_row is not None, "Sheet not found")
        if sheet_row:
            sheet_id = sheet_row[0]
            # Check header row
            cur.execute("""
                SELECT value FROM gsheet.cells
                WHERE spreadsheet_id = %s AND sheet_id = %s AND row_index = 0
                ORDER BY col_index
            """, (ss_id, sheet_id))
            headers = [r[0].lower() if r[0] else "" for r in cur.fetchall()]
            check("GSheet has Faculty Name header", any("faculty" in h and "name" in h for h in headers), f"Headers: {headers}")
            check("GSheet has Alignment Score header", any("alignment" in h or "score" in h for h in headers), f"Headers: {headers}")
            check("GSheet has Status header", any("status" in h for h in headers), f"Headers: {headers}")

            # Check data rows exist (should be 4 faculty)
            cur.execute("""
                SELECT COUNT(DISTINCT row_index) FROM gsheet.cells
                WHERE spreadsheet_id = %s AND sheet_id = %s AND row_index > 0
            """, (ss_id, sheet_id))
            data_rows = cur.fetchone()[0]
            check(f"GSheet has {FACULTY_COUNT} data rows", data_rows == FACULTY_COUNT, f"Got {data_rows} rows")

            # Check faculty names appear in the sheet
            cur.execute("""
                SELECT value FROM gsheet.cells
                WHERE spreadsheet_id = %s AND sheet_id = %s AND row_index > 0
            """, (ss_id, sheet_id))
            all_values = " ".join([r[0] for r in cur.fetchall() if r[0]]).lower()
            for fac in FACULTY:
                check(f"GSheet contains {fac}", fac.lower() in all_values, "Not found")
            # Validate Status column has only Aligned/Review Needed values
            cur.execute("""
                SELECT row_index, col_index, value FROM gsheet.cells
                WHERE spreadsheet_id = %s AND sheet_id = %s
                ORDER BY row_index, col_index
            """, (ss_id, sheet_id))
            all_cells = cur.fetchall()
            # Find Status column index from header row
            status_col = None
            for r, c, v in all_cells:
                if r == 0 and v and "status" in str(v).lower():
                    status_col = c
                    break
            if status_col is not None:
                status_values = [str(v).strip().lower() for r, c, v in all_cells if r > 0 and c == status_col and v]
                valid_set = {"aligned", "review needed", "review_needed"}
                invalid = [s for s in status_values if s not in valid_set]
                check("GSheet Status values are 'Aligned' or 'Review Needed'",
                      len(invalid) == 0, f"Invalid: {invalid}")

    # 6. Check emails sent
    # Dean email
    cur.execute("""
        SELECT id, subject, body_text FROM email.messages
        WHERE to_addr::text ILIKE '%dean@university.edu%'
        AND subject ILIKE '%Research%Teaching%Integration%Summary%'
    """)
    dean_emails = cur.fetchall()
    check("Dean email sent", len(dean_emails) >= 1, f"Found {len(dean_emails)} matching emails")
    if dean_emails:
        body = (dean_emails[0][2] or "").lower()
        check("Dean email mentions alignment", "alignment" in body or "score" in body, "Not found in body")

    # Department head emails
    cur.execute("""
        SELECT id, to_addr, subject, body_text FROM email.messages
        WHERE subject ILIKE '%Department%Research%Teaching%Alignment%Update%'
    """)
    dept_emails = cur.fetchall()
    check("Department head emails sent", len(dept_emails) >= 2, f"Found {len(dept_emails)} dept emails")

    # Check dept head recipients
    if dept_emails:
        all_recipients = " ".join([str(e[1]) for e in dept_emails]).lower()
        check("Email to biochem head", "head_biochem@university.edu" in all_recipients, f"Recipients: {all_recipients}")
        check("Email to bioinfo head", "head_bioinfo@university.edu" in all_recipients, f"Recipients: {all_recipients}")

    cur.close()
    conn.close()

    check_reverse_validation()
    # Note: Alignment_Summary.xlsx is NOT required by task.md (which only asks
    # for GSheet + Word doc + emails). The xlsx in groundtruth_workspace is a
    # reference artifact and should not be enforced on the agent.

    total = PASS_COUNT + FAIL_COUNT
    accuracy = PASS_COUNT / total * 100 if total > 0 else 0
    print(f"\nOverall: {PASS_COUNT}/{total} ({accuracy:.1f}%)")

    result = {"total_passed": PASS_COUNT, "total_checks": total, "accuracy": accuracy}
    if args.res_log_file:
        with open(args.res_log_file, "w") as f:
            json.dump(result, f, indent=2)
    sys.exit(0 if FAIL_COUNT == 0 else 1)


if __name__ == "__main__":
    main()
